"""Read a CSV or XLSX file into rows of text, with columns matched by name (plan M3.3).

Used by the import commands (``apps.core.imports``). Nothing here touches the database.

- **Formats.** CSV (UTF-8 with or without a BOM, else Windows-1252 as Excel writes it; the
  delimiter is ``;``, ``,`` or a tab, taken from the header line) and XLSX (the first
  sheet, or ``sheet``; formulas read as their last saved value; read-only streaming, with
  ``defusedxml`` guarding the XML when installed).
- **Columns.** A header matches a ``Column`` by its key or an alias, ignoring case,
  accents and the difference between space, ``_`` and ``-``. Unmapped headers are
  reported and ignored; a missing required column is an error.
- **Cells** become stripped text: numbers without a trailing ``.0``, dates as ISO dates,
  booleans as ``true`` / ``false``. Format a reference column as text in a spreadsheet,
  or a number such as ``007`` arrives as ``7``.
- A completely blank row is skipped; line numbers are the file's own (the header is line 1).
"""

import csv
import io
import re
import unicodedata
from collections.abc import Iterator, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_ROWS = 100_000
DELIMITERS = (";", ",", "\t")
CSV_ENCODINGS = ("utf-8-sig", "cp1252")
XLSX_SUFFIXES = {".xlsx", ".xlsm"}
CSV_SUFFIXES = {".csv", ".txt", ".tsv"}


class TableError(Exception):
    """The file cannot be read as a table (as opposed to a row that fails validation)."""


@dataclass(frozen=True)
class Column:
    """A column an import understands: its key, other accepted names, and whether it is needed."""

    key: str
    aliases: tuple[str, ...] = ()
    required: bool = False

    def names(self) -> set[str]:
        return {normalize_header(name) for name in (self.key, *self.aliases)}


@dataclass(frozen=True)
class Row:
    line: int
    values: dict[str, str]

    def get(self, key: str) -> str:
        return self.values.get(key, "")


@dataclass
class Table:
    source: str
    format: str
    rows: list[Row]
    mapped: dict[str, str] = field(default_factory=dict)  # column key -> header in the file
    ignored: list[str] = field(default_factory=list)  # headers no column matched
    encoding: str = ""


def normalize_header(text: str) -> str:
    folded = unicodedata.normalize("NFKD", text)
    folded = "".join(ch for ch in folded if not unicodedata.combining(ch))
    return re.sub(r"[\s_\-.]+", "_", folded.strip().casefold()).strip("_")


def _number_text(value: int | float) -> str:
    if isinstance(value, float):
        return format(Decimal(repr(value)).normalize(), "f")
    return str(value)


def _date_text(value: date) -> str:
    if isinstance(value, datetime):
        return (
            value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
        )
    return value.isoformat()


def cell_text(value: Any) -> str:
    """One cell as stripped text."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int | float):
        return _number_text(value)
    if isinstance(value, date):  # datetime is a date
        return _date_text(value)
    return str(value).strip()


def _decode(data: bytes) -> tuple[str, str]:
    for encoding in CSV_ENCODINGS:
        try:
            return data.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise TableError("The file is not UTF-8 or Windows-1252 text.")


def _delimiter(text: str, forced: str | None) -> str:
    if forced:
        return forced
    header = next(iter(text.splitlines()), "")
    return max(DELIMITERS, key=header.count)


def _csv_rows(path: Path, delimiter: str | None) -> tuple[list[list[Any]], str]:
    text, encoding = _decode(path.read_bytes())
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=_delimiter(text, delimiter))
    try:
        return list(reader), encoding
    except csv.Error as exc:
        raise TableError(f"Invalid CSV: {exc}") from exc


def _open_workbook(path: Path) -> Any:
    from openpyxl import load_workbook

    try:
        return load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # a corrupt archive raises zipfile / XML / openpyxl errors
        raise TableError(f"Not a readable XLSX file: {type(exc).__name__}") from exc


def _xlsx_rows(path: Path, sheet: str | None) -> list[list[Any]]:
    workbook = _open_workbook(path)
    try:
        if sheet is not None and sheet not in workbook.sheetnames:
            raise TableError(f"The workbook has no sheet {sheet!r} (it has {workbook.sheetnames}).")
        worksheet = workbook[sheet] if sheet is not None else workbook.worksheets[0]
        rows: list[list[Any]] = []
        for row in worksheet.iter_rows(values_only=True):
            rows.append(list(row))
            if len(rows) > MAX_ROWS + 1:  # header + MAX_ROWS; stop before memory does
                raise TableError(f"The file has more than {MAX_ROWS} rows.")
        return rows
    finally:
        workbook.close()


def _by_name(columns: Sequence[Column]) -> dict[str, Column]:
    return {name: column for column in columns for name in column.names()}


def _required_missing(columns: Sequence[Column], mapped: dict[str, str]) -> list[str]:
    return [column.key for column in columns if column.required and column.key not in mapped]


def _check_required(
    header: Sequence[Any], columns: Sequence[Column], mapped: dict[str, str]
) -> None:
    missing = _required_missing(columns, mapped)
    if not missing:
        return
    found = ", ".join(repr(title) for title in map(cell_text, header) if title) or "none"
    raise TableError(f"Required column(s) missing: {', '.join(missing)}. Found: {found}")


def _map_columns(
    header: Sequence[Any], columns: Sequence[Column]
) -> tuple[dict[int, str], dict[str, str], list[str]]:
    """``(index -> key, key -> header text, ignored headers)``; errors on a duplicate."""
    by_name = _by_name(columns)
    positions: dict[int, str] = {}
    mapped: dict[str, str] = {}
    ignored: list[str] = []
    for index, cell in enumerate(header):
        title = cell_text(cell)
        match = by_name.get(normalize_header(title)) if title else None
        if match is None:
            if title:
                ignored.append(title)
            continue
        if match.key in mapped:
            raise TableError(
                f"Column {match.key!r} appears twice: {mapped[match.key]!r}, {title!r}"
            )
        positions[index] = match.key
        mapped[match.key] = title
    _check_required(header, columns, mapped)
    return positions, mapped, ignored


def read_table(
    path: Path,
    columns: Sequence[Column],
    *,
    sheet: str | None = None,
    delimiter: str | None = None,
) -> Table:
    """Read ``path`` (CSV or XLSX by its suffix) into a ``Table`` of mapped rows."""
    _check_file(path)
    suffix = path.suffix.lower()
    encoding = ""
    if suffix in XLSX_SUFFIXES:
        raw, kind = _xlsx_rows(path, sheet), "XLSX"
    elif suffix in CSV_SUFFIXES:
        raw, encoding = _csv_rows(path, delimiter)
        kind = "CSV"
    else:
        raise TableError("Use a .csv or .xlsx file.")
    if not raw or not any(cell_text(cell) for cell in raw[0]):
        raise TableError("The file has no header row.")
    positions, mapped, ignored = _map_columns(raw[0], columns)
    rows = list(_data_rows(raw, positions))
    if len(rows) > MAX_ROWS:
        raise TableError(f"The file has more than {MAX_ROWS} rows.")
    return Table(path.name, kind, rows, mapped, ignored, encoding)


def _check_file(path: Path) -> None:
    if not path.is_file():
        raise TableError(f"{path} is not a file.")
    if path.stat().st_size > MAX_FILE_BYTES:
        raise TableError(f"The file is larger than {MAX_FILE_BYTES // (1024 * 1024)} MB.")


def _data_rows(raw: list[list[Any]], positions: dict[int, str]) -> Iterator[Row]:
    for line, cells in enumerate(raw[1:], start=2):
        values = dict.fromkeys(positions.values(), "")
        for index, key in positions.items():
            if index < len(cells):
                values[key] = cell_text(cells[index])
        if any(cell_text(cell) for cell in cells):
            yield Row(line, values)
